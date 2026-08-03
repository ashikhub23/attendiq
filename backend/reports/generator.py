import os
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def _rows_for_report(absentee_list, all_present_or_absent):
    """
    all_present_or_absent: list of dicts
      {register_number, name, status}  status in {"Present","Absent"}
    """
    return all_present_or_absent


def generate_reports(output_dir, session, student_rows, college_logo_path=None):
    """
    session: AutomationSession row (has department/section/subject/date/etc.)
    student_rows: list of dicts {register_number, name, status}
    Returns dict with pdf_path, excel_path, csv_path.
    """
    os.makedirs(output_dir, exist_ok=True)
    base_name = f"attendance_session_{session.id}"

    csv_path = os.path.join(output_dir, base_name + ".csv")
    excel_path = os.path.join(output_dir, base_name + ".xlsx")
    pdf_path = os.path.join(output_dir, base_name + ".pdf")

    _write_csv(csv_path, session, student_rows)
    _write_excel(excel_path, session, student_rows)
    _write_pdf(pdf_path, session, student_rows, college_logo_path)

    return {"csv_path": csv_path, "excel_path": excel_path, "pdf_path": pdf_path}


def _write_csv(path, session, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Register Number", "Student Name", "Status", "Date",
                          "Department", "Section", "Subject"])
        for r in rows:
            writer.writerow([
                r["register_number"], r.get("name") or "", r["status"],
                session.date, session.department, session.section, session.subject
            ])
        writer.writerow([])
        writer.writerow(["Present", session.total_present])
        writer.writerow(["Absent", session.total_absent])
        writer.writerow(["Automation Duration (sec)", session.duration_seconds])
        writer.writerow(["Automation Status", session.status])


def _write_excel(path, session, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = ["Register Number", "Student Name", "Status", "Date", "Department", "Section", "Subject"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([
            r["register_number"], r.get("name") or "", r["status"],
            session.date, session.department, session.section, session.subject
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(12, max_len + 2)

    summary_row = ws.max_row + 2
    ws.cell(row=summary_row, column=1, value="Present").font = Font(bold=True)
    ws.cell(row=summary_row, column=2, value=session.total_present)
    ws.cell(row=summary_row + 1, column=1, value="Absent").font = Font(bold=True)
    ws.cell(row=summary_row + 1, column=2, value=session.total_absent)
    ws.cell(row=summary_row + 2, column=1, value="Automation Status").font = Font(bold=True)
    ws.cell(row=summary_row + 2, column=2, value=session.status)

    wb.save(path)


def _write_pdf(path, session, rows, logo_path=None):
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle", parent=styles["Title"], textColor=colors.HexColor("#1E3A8A")
    )
    elements = []

    if logo_path and os.path.exists(logo_path):
        elements.append(Image(logo_path, width=3 * cm, height=3 * cm))
        elements.append(Spacer(1, 8))

    elements.append(Paragraph("Attendance Report", title_style))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(
        f"Department: {session.department or '-'} &nbsp;&nbsp; "
        f"Section: {session.section or '-'} &nbsp;&nbsp; "
        f"Subject: {session.subject or '-'} &nbsp;&nbsp; "
        f"Date: {session.date or '-'}", styles["Normal"]
    ))
    elements.append(Paragraph(
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]
    ))
    elements.append(Spacer(1, 12))

    table_data = [["Register Number", "Student Name", "Status"]]
    for r in rows:
        table_data.append([r["register_number"], r.get("name") or "-", r["status"]])

    table = Table(table_data, colWidths=[5.5 * cm, 7 * cm, 3.5 * cm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 16))

    summary_data = [
        ["Present", str(session.total_present)],
        ["Absent", str(session.total_absent)],
        ["Automation Duration", f"{session.duration_seconds} sec"],
        ["Automation Status", session.status.upper()],
    ]
    summary_table = Table(summary_data, colWidths=[6 * cm, 6 * cm])
    summary_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
    ]))
    elements.append(summary_table)

    doc.build(elements)
